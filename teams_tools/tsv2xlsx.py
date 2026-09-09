import csv
import io
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.utils import get_column_letter
from _paths import tsv_path, xlsx_path

# Lookup pages (Speed, Strategy, ...) are authored as <dst-dir>/<name>.tsv next
# to the master teams.tsv. Files may have a leading block of '#' comment lines;
# data fields may contain embedded newlines inside double quotes. Optional:
# absent file -> that sheet is skipped.
def lookup_rows(dst, filename):
    p = Path(dst).resolve().parent / filename
    if not p.exists():
        return None
    text = open(p, encoding="utf-8").read()
    lines = text.split("\n")
    i = 0
    while i < len(lines) and lines[i].lstrip().startswith("#"):
        i += 1
    rows = list(csv.reader(io.StringIO("\n".join(lines[i:])), delimiter="\t"))
    return rows[0], rows[1:]

def pad_rows(header, data):
    out = []
    for r in data:
        out.append(r[: len(header)] + [""] * max(0, len(header) - len(r)))
    return out

src = tsv_path(sys.argv[1] if len(sys.argv) > 1 else None)
dst = xlsx_path(sys.argv[2] if len(sys.argv) > 2 else None)

lines = open(src, encoding="utf-8").read().splitlines()
table = [ln for ln in lines if not ln.startswith("#")]
rows = list(csv.reader(table, delimiter="\t"))
header, data = rows[0], rows[1:]

# resolve columns by header name -> robust to column order
def col(name):
    return header.index(name)

assert all(len(r) == len(header) for r in data), "column count mismatch"

C = {name: col(name) for name in header}
I_NAME, I_TIER, I_TAGS = C["Pokemon"], C["Tier"], C["Tags"]

TIER_ORDER = ["AG", "Uber", "OU", "UUBL", "UU", "RUBL", "RU", "below RU", "MISC"]

def tier_rank(tier):
    return TIER_ORDER.index(tier) if tier in TIER_ORDER else len(TIER_ORDER)

def sort_rows(rs):
    return sorted(
        rs,
        key=lambda r: (
            0 if "Mega" in r[I_NAME] else 1,
            tier_rank(r[I_TIER]),
            r[I_NAME].lower(),
        ),
    )

# combined tier tabs
def sheets_for(data):
    megas = lambda r: "Mega" in r[I_NAME]
    tiers_of = lambda *ts: (lambda r: r[I_TIER] in ts)
    return [
        ("Mega", megas),
        ("Uber", tiers_of("AG", "Uber")),
        ("OU", tiers_of("OU", "UUBL")),
        ("UU", tiers_of("UU", "RUBL")),
        ("RU", tiers_of("RU", "below RU")),
        ("MISC", lambda r: False),
    ]

sheets = sheets_for(data)

TAG_COLORS = {
    "SetUp": "7030A0",
    "Fast": "C00000",
    "Breaker": "BF8F00",
    "Utility": "00B050",
    "Wall": "1F6E9C",
}
SEP_COLOR = "808080"

TIER_FILLS = {
    "AG": "F8CBAD",
    "Uber": "F4CCCC",
    "OU": "D9EAD3",
    "UUBL": "D0E0E3",
    "UU": "CFE2F3",
    "RUBL": "D9D2E9",
    "RU": "FFF2CC",
    "below RU": "E7E6E6",
    "MISC": "F2F2F2",
}

def tag_richtext(value):
    if not value:
        return ""
    blocks = []
    for i, token in enumerate(value.split("/")):
        token = token.strip()
        if i > 0:
            blocks.append(TextBlock(InlineFont(color=SEP_COLOR), " / "))
        blocks.append(TextBlock(InlineFont(color=TAG_COLORS.get(token, "000000")), token))
    return CellRichText(blocks)

widths = [22, 14, 32, 10, 26, 34, 6, 6, 6, 6, 6, 6]

wb = Workbook()
wb.remove(wb.active)

lookup = lookup_rows(dst, "speed.tsv")
if lookup is not None:
    ws = wb.create_sheet("Speed")
    ws.append(lookup[0])
    for r in pad_rows(lookup[0], lookup[1]):
        row = [int(v) if str(v).isdigit() else v for v in r]
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for i, w in enumerate([16, 16, 16, 16], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

lookup = lookup_rows(dst, "strategy.tsv")
if lookup is not None:
    ws = wb.create_sheet("Strategy")
    ws.append(lookup[0])
    for r in pad_rows(lookup[0], lookup[1]):
        ws.append(r)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    for i, w in enumerate([24, 10, 30, 30, 90, 75], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

for name, pred in sheets:
    ws = wb.create_sheet(name)
    ws.append(header)
    for r in sort_rows([r for r in data if pred(r)]):
        ws.append(r)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        tier_cell = row[I_TIER]
        if tier_cell.value in TIER_FILLS:
            tier_cell.fill = PatternFill("solid", fgColor=TIER_FILLS[tier_cell.value])
        tag_cell = row[I_TAGS]
        tag_cell.value = tag_richtext(tag_cell.value)
        # numeric stats
        for stat in ("HP", "ATK", "DEF", "SPA", "SPD", "SPE"):
            c = header.index(stat)
            cell = row[c]
            if cell.value is not None and str(cell.value).isdigit():
                cell.value = int(cell.value)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate(widths[: len(header)], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

wb.save(dst)
print(f"saved {dst}")
for name, pred in sheets:
    print(f"  {name}: {sum(1 for r in data if pred(r))}")
